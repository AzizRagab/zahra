"""
RAG HACK - Full Content Extraction Script
Extracts all text content (books, images, subtitles, docs) from D:\RAG HACK
into a single organized folder ready for copying: D:\RAG_Extracted
"""
import os
import sys
import shutil
import traceback
from pathlib import Path
from datetime import datetime

SRC = Path(r"D:\RAG HACK")
OUT = Path(r"D:\RAG_Extracted")

# ---------- Output structure ----------
BOOKS_TXT = OUT / "01_Books_Text"          # Extracted text from PDFs
DOCS_TXT = OUT / "02_Docs_Text"            # DOCX/PPTX/TXT/MD/HTML text
SUBS_DIR = OUT / "03_Subtitles"            # VTT files (copied)
IMAGES_DIR = OUT / "04_Images"             # Images (copied)
OCR_DIR = OUT / "04_Images" / "_OCR_Text"  # OCR text from images
ORIGINALS = OUT / "05_Original_Files"      # Original PDF/EPUB/DJVU copies
VIDEOS_LIST = OUT / "06_Videos_List.txt"   # List of all videos
INDEX_FILE = OUT / "00_INDEX.txt"          # Master index
REPORT_FILE = OUT / "00_EXTRACTION_REPORT.txt"

# ---------- Stats ----------
stats = {
    "pdf_total": 0, "pdf_ok": 0, "pdf_failed": 0,
    "docx_total": 0, "docx_ok": 0,
    "pptx_total": 0, "pptx_ok": 0,
    "vtt_total": 0, "vtt_copied": 0,
    "img_total": 0, "img_copied": 0, "img_ocr_done": 0,
    "txt_total": 0,
    "videos_total": 0,
    "originals_copied": 0,
    "errors": [],
}

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff", ".jfif"}
DOC_TEXT_EXTS = {".txt", ".md", ".html", ".htm", ".xml", ".yml", ".yaml", ".json", ".csv", ".log"}
ORIGINAL_EXTS = {".pdf", ".epub", ".djvu"}
SUBTITLE_EXTS = {".vtt", ".srt"}
VIDEO_EXTS = {".mp4", ".avi", ".mkv", ".mov", ".wmv", ".flv", ".webm", ".m4v"}


def log(msg):
    print(msg, flush=True)


def safe_name(path: Path, max_len=120) -> str:
    name = path.name
    if len(name) > max_len:
        stem = path.stem[: max_len - 10]
        name = stem + path.suffix
    return name


def ensure_out_dirs():
    for d in [BOOKS_TXT, DOCS_TXT, SUBS_DIR, IMAGES_DIR, OCR_DIR, ORIGINALS]:
        d.mkdir(parents=True, exist_ok=True)


# =====================================================================
# PDF text extraction (pypdf)
# =====================================================================
def extract_pdf_text(pdf_path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(pdf_path))
    parts = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
            if t.strip():
                parts.append(t)
        except Exception:
            continue
    return "\n\n--- PAGE BREAK ---\n\n".join(parts)


def process_pdfs():
    for pdf in sorted(SRC.rglob("*.pdf")):
        stats["pdf_total"] += 1
        try:
            rel = pdf.relative_to(SRC)
            out_rel = rel.with_suffix(".txt")
            # flatten directory structure to avoid too many nested dirs
            parts = list(out_rel.parts)
            if len(parts) > 3:
                folder_key = "_".join(parts[0])
                out_file = BOOKS_TXT / f"{folder_key}__{'_'.join(parts[1:])}"
            else:
                out_file = BOOKS_TXT / "_".join(parts)
            out_file.parent.mkdir(parents=True, exist_ok=True)
            text = extract_pdf_text(pdf)
            if text.strip():
                header = f"SOURCE: {pdf}\nEXTRACTED: {datetime.now().isoformat()}\n\n"
                out_file.write_text(header + text, encoding="utf-8", errors="replace")
                stats["pdf_ok"] += 1
                log(f"  [PDF OK] {pdf.name} -> {len(text):,} chars")
            else:
                stats["pdf_failed"] += 1
                stats["errors"].append(f"EMPTY TEXT: {pdf}")
                log(f"  [PDF EMPTY] {pdf.name}")
        except Exception as e:
            stats["pdf_failed"] += 1
            stats["errors"].append(f"PDF FAIL: {pdf} :: {e}")
            log(f"  [PDF FAIL] {pdf.name} :: {e}")


# =====================================================================
# DOCX text extraction (python-docx)
# =====================================================================
def process_docx():
    for docx in sorted(SRC.rglob("*.docx")):
        stats["docx_total"] += 1
        try:
            from docx import Document
            doc = Document(str(docx))
            parts = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
            rel = docx.relative_to(SRC)
            out_file = DOCS_TXT / "_".join(rel.parts[:-1] + (rel.stem + ".txt",))
            out_file.parent.mkdir(parents=True, exist_ok=True)
            header = f"SOURCE: {docx}\nEXTRACTED: {datetime.now().isoformat()}\n\n"
            out_file.write_text(header + text, encoding="utf-8", errors="replace")
            stats["docx_ok"] += 1
            log(f"  [DOCX OK] {docx.name} -> {len(text):,} chars")
        except Exception as e:
            stats["errors"].append(f"DOCX FAIL: {docx} :: {e}")
            log(f"  [DOCX FAIL] {docx.name} :: {e}")


# =====================================================================
# PPTX text extraction (python-pptx)
# =====================================================================
def process_pptx():
    for pptx in sorted(SRC.rglob("*.pptx")):
        stats["pptx_total"] += 1
        try:
            from pptx import Presentation
            prs = Presentation(str(pptx))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                slide_parts = [f"=== SLIDE {i} ==="]
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = "".join(run.text for run in para.runs)
                            if t.strip():
                                slide_parts.append(t)
                    if shape.shape_type == 19:  # table
                        try:
                            for row in shape.table.rows:
                                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                                if cells:
                                    slide_parts.append(" | ".join(cells))
                        except Exception:
                            pass
                if len(slide_parts) > 1:
                    parts.append("\n".join(slide_parts))
            text = "\n\n".join(parts)
            rel = pptx.relative_to(SRC)
            out_file = DOCS_TXT / "_".join(rel.parts[:-1] + (rel.stem + ".txt",))
            out_file.parent.mkdir(parents=True, exist_ok=True)
            header = f"SOURCE: {pptx}\nEXTRACTED: {datetime.now().isoformat()}\n\n"
            out_file.write_text(header + text, encoding="utf-8", errors="replace")
            stats["pptx_ok"] += 1
            log(f"  [PPTX OK] {pptx.name} -> {len(text):,} chars")
        except Exception as e:
            stats["errors"].append(f"PPTX FAIL: {pptx} :: {e}")
            log(f"  [PPTX FAIL] {pptx.name} :: {e}")


# =====================================================================
# Plain text files (txt / md / html etc.)
# =====================================================================
def process_text_files():
    for f in sorted(SRC.rglob("*")):
        if not f.is_file():
            continue
        if f.suffix.lower() not in DOC_TEXT_EXTS:
            continue
        # skip our own workspace files
        if "RAG_Extracted" in str(f):
            continue
        stats["txt_total"] += 1
        try:
            rel = f.relative_to(SRC)
            out_file = DOCS_TXT / "Text_Files" / "_".join(rel.parts)
            out_file.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, out_file)
        except Exception:
            pass


# =====================================================================
# Subtitles (VTT/SRT) - copy directly
# =====================================================================
def process_subtitles():
    for f in sorted(SRC.rglob("*")):
        if not f.is_file():
            continue
        if f.suffix.lower() not in SUBTITLE_EXTS:
            continue
        stats["vtt_total"] += 1
        try:
            rel = f.relative_to(SRC)
            out_file = SUBS_DIR / "_".join(rel.parts)
            out_file.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, out_file)
            stats["vtt_copied"] += 1
        except Exception as e:
            stats["errors"].append(f"VTT COPY FAIL: {f} :: {e}")


# =====================================================================
# Images - copy + OCR text extraction
# =====================================================================
def process_images(do_ocr=True):
    from PIL import Image
    try:
        from rapidocr_onnxruntime import RapidOCR
        ocr_engine = RapidOCR() if do_ocr else None
        log("  [OCR] Engine loaded.")
    except Exception as e:
        ocr_engine = None
        log(f"  [OCR] Engine unavailable: {e}")
        stats["errors"].append(f"OCR engine unavailable: {e}")

    for img in sorted(SRC.rglob("*")):
        if not img.is_file():
            continue
        if img.suffix.lower() not in IMAGE_EXTS:
            continue
        stats["img_total"] += 1
        try:
            rel = img.relative_to(SRC)
            out_img = IMAGES_DIR / "_".join(rel.parts)
            out_img.parent.mkdir(parents=True, exist_ok=True)
            # downscale if huge (>12MP) to save space
            try:
                with Image.open(img) as im:
                    w, h = im.size
                    if w * h > 12_000_000:
                        ratio = (12_000_000 / (w * h)) ** 0.5
                        im = im.resize((int(w * ratio), int(h * ratio)), Image.LANCZOS)
                        # convert RGBA/P to RGB if saving as jpg
                        ext = out_img.suffix.lower()
                        if ext in (".jpg", ".jpeg") and im.mode != "RGB":
                            im = im.convert("RGB")
                        if ext in (".png", ".webp") and im.mode == "P":
                            im = im.convert("RGBA" if "transparency" in im.info else "RGB")
                        im.save(out_img, optimize=True)
                    else:
                        shutil.copy2(img, out_img)
            except Exception:
                shutil.copy2(img, out_img)
            stats["img_copied"] += 1

            # ---- OCR ----
            if ocr_engine:
                try:
                    # verify image is valid before OCR
                    with Image.open(img) as im:
                        im.verify()
                    result, _ = ocr_engine(str(img))
                    if result:
                        lines = [item[1] for item in result]
                        text = "\n".join(lines)
                        if text.strip():
                            ocr_file = OCR_DIR / ("_".join(rel.parts[:-1]) + "__" + rel.stem + ".txt")
                            ocr_file.parent.mkdir(parents=True, exist_ok=True)
                            header = f"SOURCE IMAGE: {img}\n\n"
                            ocr_file.write_text(header + text, encoding="utf-8", errors="replace")
                            stats["img_ocr_done"] += 1
                            log(f"  [IMG OCR OK] {img.name} -> {len(text):,} chars")
                        else:
                            log(f"  [IMG OCR EMPTY] {img.name}")
                    else:
                        log(f"  [IMG OCR EMPTY] {img.name}")
                except Exception as e:
                    stats["errors"].append(f"IMG OCR FAIL: {img} :: {e}")
                    log(f"  [IMG OCR FAIL] {img.name} :: {e}")
            else:
                log(f"  [IMG COPIED] {img.name} (no OCR)")
        except Exception as e:
            stats["errors"].append(f"IMG COPY FAIL: {img} :: {e}")
            log(f"  [IMG COPY FAIL] {img.name} :: {e}")


# =====================================================================
# Original books (PDF/EPUB/DJVU) - copied as-is for reference
# =====================================================================
def process_originals():
    for f in sorted(SRC.rglob("*")):
        if not f.is_file():
            continue
        if f.suffix.lower() not in ORIGINAL_EXTS:
            continue
        try:
            rel = f.relative_to(SRC)
            out_file = ORIGINALS / "_".join(rel.parts)
            out_file.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(f, out_file)
            stats["originals_copied"] += 1
        except Exception as e:
            stats["errors"].append(f"ORIGINAL COPY FAIL: {f} :: {e}")


# =====================================================================
# Videos - list only (too large to copy)
# =====================================================================
def process_videos():
    lines = []
    total_size = 0
    for f in sorted(SRC.rglob("*")):
        if not f.is_file():
            continue
        if f.suffix.lower() not in VIDEO_EXTS:
            continue
        stats["videos_total"] += 1
        size_mb = f.stat().st_size / (1024 * 1024)
        total_size += size_mb
        lines.append(f"{size_mb/1024:8.2f} GB\t{f}")
    header = (
        f"RAG HACK - ALL VIDEO FILES\n"
        f"Generated: {datetime.now().isoformat()}\n"
        f"Total Videos: {stats['videos_total']}\n"
        f"Total Video Size: {total_size/1024:.2f} GB\n"
        f"{'='*80}\n\n"
    )
    VIDEOS_LIST.write_text(header + "\n".join(lines), encoding="utf-8")


# =====================================================================
# Master index
# =====================================================================
def write_index():
    lines = []
    lines.append("=" * 90)
    lines.append("RAG HACK - FULL CONTENT EXTRACTION INDEX")
    lines.append("Generated: " + datetime.now().isoformat())
    lines.append("Source: " + str(SRC))
    lines.append("=" * 90)
    lines.append("")
    lines.append("CONTENTS:")
    lines.append(f"  01_Books_Text        - Extracted text from {stats['pdf_ok']} PDF books")
    lines.append(f"  02_Docs_Text         - Extracted text from DOCX/PPTX/TXT/MD/HTML files")
    lines.append(f"  03_Subtitles         - {stats['vtt_copied']} video subtitle files (VTT/SRT)")
    lines.append(f"  04_Images            - {stats['img_copied']} images + OCR text in _OCR_Text")
    lines.append(f"  05_Original_Files    - {stats['originals_copied']} original book files (PDF/EPUB/DJVU)")
    lines.append(f"  06_Videos_List.txt   - {stats['videos_total']} videos (paths listed, files too large to copy)")
    lines.append("")
    lines.append("=" * 90)
    lines.append("SAMPLE FILE LISTING")
    lines.append("=" * 90)
    lines.append("")

    # list top-level contents of each folder
    for folder in [BOOKS_TXT, DOCS_TXT, SUBS_DIR, IMAGES_DIR, OCR_DIR, ORIGINALS]:
        if folder.exists():
            items = sorted(folder.rglob("*")) if folder != BOOKS_TXT else sorted(folder.glob("*.txt"))
            lines.append(f"\n--- {folder.name} ({len(list(folder.rglob('*')) if folder != BOOKS_TXT else items)} items) ---")
            shown = 0
            for it in items[:50]:
                if it.is_file():
                    size = it.stat().st_size
                    lines.append(f"   {size:>12,}  {it}")
                    shown += 1
            if shown >= 50:
                lines.append(f"   ... and more ({max(0, len(items)-50)} additional files)")
    INDEX_FILE.write_text("\n".join(lines), encoding="utf-8")


# =====================================================================
# Report
# =====================================================================
def write_report():
    elapsed = datetime.now() - start_time
    lines = []
    lines.append("=" * 90)
    lines.append("RAG HACK - EXTRACTION REPORT")
    lines.append("=" * 90)
    lines.append(f"Started:    {start_time.isoformat()}")
    lines.append(f"Finished:   {datetime.now().isoformat()}")
    lines.append(f"Elapsed:    {elapsed}")
    lines.append("")
    lines.append(f"PDF books found:        {stats['pdf_total']}")
    lines.append(f"  - extracted OK:       {stats['pdf_ok']}")
    lines.append(f"  - empty/failed:       {stats['pdf_failed']}")
    lines.append(f"DOCX files:             {stats['docx_total']} (OK: {stats['docx_ok']})")
    lines.append(f"PPTX files:             {stats['pptx_total']} (OK: {stats['pptx_ok']})")
    lines.append(f"Text files processed:   {stats['txt_total']}")
    lines.append(f"Subtitle files copied:  {stats['vtt_copied']} / {stats['vtt_total']}")
    lines.append(f"Images copied:          {stats['img_copied']} / {stats['img_total']}")
    lines.append(f"Images OCR'd:           {stats['img_ocr_done']}")
    lines.append(f"Original books copied:  {stats['originals_copied']}")
    lines.append(f"Videos listed:          {stats['videos_total']}")
    lines.append("")
    if stats["errors"]:
        lines.append(f"ERRORS ({len(stats['errors'])}):")
        for e in stats["errors"][:100]:
            lines.append(f"  - {e}")
        if len(stats["errors"]) > 100:
            lines.append(f"  ... and {len(stats['errors'])-100} more")
    REPORT_FILE.write_text("\n".join(lines), encoding="utf-8")


# =====================================================================
# MAIN
# =====================================================================
if __name__ == "__main__":
    start_time = datetime.now()
    if not SRC.exists():
        print(f"ERROR: Source folder not found: {SRC}")
        sys.exit(1)

    do_ocr = "--no-ocr" not in sys.argv
    if not do_ocr:
        print("[MODE] OCR disabled (use --no-ocr flag detected)")
    else:
        print("[MODE] OCR enabled (add --no-ocr to skip)")

    print(f"[*] Source: {SRC}")
    print(f"[*] Target: {OUT}")
    print("[*] Creating output structure...")
    ensure_out_dirs()

    print("\n[1/7] Extracting PDF book texts...")
    process_pdfs()

    print("\n[2/7] Extracting DOCX texts...")
    process_docx()

    print("\n[3/7] Extracting PPTX texts...")
    process_pptx()

    print("\n[4/7] Copying text files...")
    process_text_files()

    print("\n[5/7] Copying subtitles (VTT/SRT)...")
    process_subtitles()

    print("\n[6/7] Copying images + OCR text extraction...")
    process_images(do_ocr=do_ocr)

    print("\n[7/7] Copying original books + writing index/report...")
    process_originals()
    process_videos()
    write_index()
    write_report()

    print("\n" + "=" * 60)
    print("EXTRACTION COMPLETE")
    print("=" * 60)
    print(f"Output folder: {OUT}")
    print(f"  PDFs extracted:   {stats['pdf_ok']}/{stats['pdf_total']}")
    print(f"  DOCX extracted:   {stats['docx_ok']}/{stats['docx_total']}")
    print(f"  PPTX extracted:   {stats['pptx_ok']}/{stats['pptx_total']}")
    print(f"  Images copied:    {stats['img_copied']}/{stats['img_total']}")
    print(f"  Images OCR'd:     {stats['img_ocr_done']}")
    print(f"  Subtitles:        {stats['vtt_copied']}")
    print(f"  Originals copied: {stats['originals_copied']}")
    print(f"  Videos listed:    {stats['videos_total']}")
    if stats["errors"]:
